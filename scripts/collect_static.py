"""Generate static knowledge snapshots. Runs locally or in GitHub Actions."""
import csv
import io
import json
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

import requests
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts' / 'kpl'))
import kpl_official as kpl

OUT = ROOT / 'public' / 'rag'
COLS = ['主题/对象', '知识类型', '知识内容', '解说具体场景', '时效性', '可信度', '来源名称', '来源链接', '备注']


def generate():
    config = json.loads((ROOT / 'scripts' / 'rag-config.json').read_text(encoding='utf-8'))
    now = datetime.now(timezone(timedelta(hours=8)))
    data = json.loads((OUT / 'seed.json').read_text(encoding='utf-8'))
    # Keep the previous successful live snapshot if today's upstream fails.
    try:
        response = requests.get('https://asta.net.cn/rag/data.json', timeout=15)
        response.raise_for_status()
        previous = response.json()
        if isinstance(previous.get('rows'), list) and previous.get('last_success'):
            data = previous
    except (requests.RequestException, ValueError, AttributeError):
        pass
    archive = json.loads((OUT / 'seed.json').read_text(encoding='utf-8'))['rows']
    data['last_attempt'] = now.isoformat()
    data['season'] = config['season']
    try:
        schedules = sorted(kpl.list_schedules(config['season']), key=lambda s: s['timestamp'], reverse=True)
        done = [s for s in schedules if s['status'] == 4]
        if not done:
            raise RuntimeError('官方接口未返回已完成赛程，保留上次结果')
        live = []
        for s in done:
            a, b = s['team_a_name'], s['team_b_name']
            live.append(dict(zip(COLS, [f'{a} vs {b}', '战队赛果',
                f"{s['date']} {s['stage_name']}：{a} {s['team_a_score']}:{s['team_b_score']} {b}。",
                '赛前介绍、交手回顾、近期战绩', s['date'], '高', 'KPL官方赛程接口',
                f"{kpl.KPL_BASE}/getScheduleList?seasonid={config['season']}",
                f"赛季{config['season']}；对阵{s['scheduleid']}"])))
        detail_ok = 0
        warnings = []
        for s in done[:config['detail_limit']]:
            try:
                detail = kpl.schedule_detail(config['season'], s['scheduleid'])
                if not detail['rounds']:
                    raise RuntimeError('没有逐局详情')
                detail_ok += 1
                pools = {}
                for rd in detail['rounds']:
                    for p in rd['players']:
                        if p.get('playerid') and p.get('hero_name'):
                            pools.setdefault(p['playerid'], []).append(p['hero_name'])
                for pid, heroes in pools.items():
                    p = detail['players'].get(pid, {})
                    name = p.get('player_short') or p.get('player_name') or str(pid)
                    live.append(dict(zip(COLS, [name, '选手英雄池',
                        f"{name}在{s['date']} {s['team_a_name']} vs {s['team_b_name']}使用：{'、'.join(heroes)}。",
                        'BP环节、英雄池与对位分析', s['date'], '高', 'KPL官方对局详情',
                        f"{kpl.KPL_BASE}/getScheduleDetail?seasonid={config['season']}&scheduleid={s['scheduleid']}",
                        f"赛季{config['season']}；对阵{s['scheduleid']}"])))
            except Exception as exc:
                warnings.append(f"对阵{s['scheduleid']}详情未更新：{exc}")
        for row in live:
            row['数据批次'] = '自动采集'
        data.update(rows=archive + live, last_success=now.isoformat(), status='success' if not warnings else 'partial',
                    message=f'更新 {len(done)} 场赛果，{detail_ok} 场详情。' + (' 部分详情不可用。' if warnings else ''), warnings=warnings)
    except Exception as exc:
        data.update(status='failed', message=f'本次自动采集失败，继续展示已有数据：{exc}')
        print(data['message'], file=sys.stderr)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / 'data.json').write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
    buf = io.StringIO(newline='')
    writer = csv.DictWriter(buf, fieldnames=COLS, extrasaction='ignore')
    writer.writeheader()
    writer.writerows(data['rows'])
    (OUT / 'knowledge.csv').write_bytes(buf.getvalue().encode('utf-8-sig'))
    wb = Workbook()
    ws = wb.active
    ws.title = '知识库'
    ws.append(COLS)
    for row in data['rows']:
        # Treat upstream text as text even if it starts with an Excel formula marker.
        ws.append([str(row.get(c, '')) for c in COLS])
        for cell in ws[ws.max_row]:
            cell.data_type = 's'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(OUT / 'knowledge.xlsx')
    print(f"Snapshot: {len(data['rows'])} rows; {data['status']}")


if __name__ == '__main__':
    generate()
